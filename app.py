import os,io,math,secrets,hashlib
from datetime import datetime,timezone,timedelta
from functools import wraps
import requests
from flask import Flask,render_template,request,jsonify,session,redirect,url_for,send_file
from dotenv import load_dotenv
from openpyxl import load_workbook
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import generate_password_hash,check_password_hash
load_dotenv()
app=Flask(__name__)
app.config.update(SECRET_KEY=os.getenv('FLASK_SECRET') or 'change-this-secret',SESSION_COOKIE_NAME='corva_admin_session_v2',SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SECURE=True,SESSION_COOKIE_SAMESITE='Lax')
app.wsgi_app=ProxyFix(app.wsgi_app,x_proto=1,x_host=1)
SUPABASE_URL=os.getenv('SUPABASE_URL','').rstrip('/')
SUPABASE_KEY=os.getenv('SUPABASE_SERVICE_KEY','')
ADMIN_USERNAME=os.getenv('ADMIN_USERNAME','CORVA_ADMIN')
ADMIN_PASSWORD=os.getenv('ADMIN_PASSWORD','change-me')
ADMIN_EMAIL=os.getenv('ADMIN_EMAIL','').strip().lower()
RESEND_API_KEY=os.getenv('RESEND_API_KEY','')
RESEND_FROM_EMAIL=os.getenv('RESEND_FROM_EMAIL','onboarding@resend.dev')

def hdr(): return {'apikey':SUPABASE_KEY,'Authorization':f'Bearer {SUPABASE_KEY}','Content-Type':'application/json'}
def get(t,p=None):
 r=requests.get(f'{SUPABASE_URL}/rest/v1/{t}',headers=hdr(),params=p or {},timeout=20);r.raise_for_status();return r.json()
def insert(t,d):
 r=requests.post(f'{SUPABASE_URL}/rest/v1/{t}',headers={**hdr(),'Prefer':'return=representation'},json=d,timeout=20);r.raise_for_status();return r.json()
def update_where(t,params,d):
 r=requests.patch(f'{SUPABASE_URL}/rest/v1/{t}',headers={**hdr(),'Prefer':'return=representation'},params=params,json=d,timeout=20);r.raise_for_status();return r.json()
def delete_where(t,params):
 r=requests.delete(f'{SUPABASE_URL}/rest/v1/{t}',headers=hdr(),params=params,timeout=20);r.raise_for_status()
def balance(x): return max(0,float(x.get('total_qty') or 0)-float(x.get('received') or 0))
def buildable(rows):
 vals=[]
 for x in rows:
  bom=float(x.get('bom_qty') or 0);received=float(x.get('received') or 0);sets=float(x.get('no_of_sets') or 0)
  if bom>0:
   possible=math.floor(received/bom)
   if sets>0: possible=min(possible,math.floor(sets))
   vals.append(possible)
 return min(vals) if vals else 0

def admin_account():
 rows=get('admin_accounts',{'select':'*','username':f'eq.{ADMIN_USERNAME}','limit':'1'})
 if rows:return rows[0]
 if not ADMIN_EMAIL:return None
 try:return insert('admin_accounts',{'username':ADMIN_USERNAME,'email':ADMIN_EMAIL,'password_hash':generate_password_hash(ADMIN_PASSWORD)})[0]
 except Exception:return None

def require_admin(f):
 @wraps(f)
 def w(*a,**k):
  if session.get('role')!='admin': return jsonify(error='Admin access required'),403
  return f(*a,**k)
 return w

@app.after_request
def security_headers(r):
 r.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0';r.headers['Pragma']='no-cache';return r
@app.get('/')
def home(): return render_template('index.html',role=session.get('role','user'),username=session.get('user','User'))
@app.get('/login')
def login(): return redirect(url_for('home'))
@app.post('/api/login')
def api_login():
 d=request.get_json(silent=True) or {};u=str(d.get('username','')).strip();p=str(d.get('password',''))
 if u!=ADMIN_USERNAME:return jsonify(ok=False,error='Invalid admin credentials'),401
 try:a=admin_account();valid=check_password_hash(a['password_hash'],p) if a else (p==ADMIN_PASSWORD)
 except Exception:valid=(p==ADMIN_PASSWORD)
 if valid: session.clear();session['user']=u;session['role']='admin';return jsonify(ok=True,role='admin')
 return jsonify(ok=False,error='Invalid admin credentials'),401
@app.post('/api/forgot/request')
def forgot_request():
 d=request.get_json(silent=True) or {};u=str(d.get('username','')).strip();email=str(d.get('email','')).strip().lower()
 if u!=ADMIN_USERNAME or not email:return jsonify(error='Enter Admin ID and registered email.'),400
 a=admin_account()
 if not a or email!=str(a.get('email','')).strip().lower():return jsonify(error='Admin ID or registered email is incorrect.'),401
 if not RESEND_API_KEY:return jsonify(error='Email reset is not configured on the server yet.'),503
 otp=f'{secrets.randbelow(1000000):06d}';otp_hash=hashlib.sha256(otp.encode()).hexdigest();expires=(datetime.now(timezone.utc)+timedelta(minutes=10)).isoformat()
 try:
  delete_where('admin_password_reset',{'username':f'eq.{ADMIN_USERNAME}'})
  insert('admin_password_reset',{'username':ADMIN_USERNAME,'otp_hash':otp_hash,'expires_at':expires,'attempts':0})
  r=requests.post('https://api.resend.com/emails',headers={'Authorization':f'Bearer {RESEND_API_KEY}','Content-Type':'application/json'},json={'from':RESEND_FROM_EMAIL,'to':[email],'subject':'CORVA Admin Password Reset OTP','text':f'Your CORVA Fixture Tracker password reset OTP is {otp}. It expires in 10 minutes. If you did not request this, ignore this email.'},timeout=20);r.raise_for_status();return jsonify(ok=True,message='OTP sent to your registered email.')
 except Exception:return jsonify(error='Could not send OTP. Check Resend configuration.'),502
@app.post('/api/forgot/reset')
def forgot_reset():
 d=request.get_json(silent=True) or {};u=str(d.get('username','')).strip();email=str(d.get('email','')).strip().lower();otp=str(d.get('otp','')).strip();pw=str(d.get('new_password',''))
 if u!=ADMIN_USERNAME or len(otp)!=6 or len(pw)<8:return jsonify(error='Enter valid reset details.'),400
 rows=get('admin_password_reset',{'select':'*','username':f'eq.{ADMIN_USERNAME}','order':'id.desc','limit':'1'})
 if not rows:return jsonify(error='No active OTP. Request a new one.'),400
 x=rows[0]
 try:expired=datetime.fromisoformat(str(x['expires_at']).replace('Z','+00:00'))<=datetime.now(timezone.utc)
 except Exception:expired=True
 if expired:return jsonify(error='OTP expired. Request a new one.'),400
 if int(x.get('attempts') or 0)>=5:return jsonify(error='Too many attempts. Request a new OTP.'),429
 if hashlib.sha256(otp.encode()).hexdigest()!=str(x.get('otp_hash','')):
  update_where('admin_password_reset',{'id':f"eq.{x['id']}"},{'attempts':int(x.get('attempts') or 0)+1});return jsonify(error='Invalid OTP.'),401
 a=admin_account()
 if not a or email!=str(a.get('email','')).strip().lower():return jsonify(error='Registered email is incorrect.'),401
 update_where('admin_accounts',{'username':f'eq.{ADMIN_USERNAME}'},{'password_hash':generate_password_hash(pw),'updated_at':datetime.now(timezone.utc).isoformat()});delete_where('admin_password_reset',{'id':f"eq.{x['id']}"});return jsonify(ok=True,message='Password reset successfully.')
@app.post('/api/admin/change-password')
@require_admin
def change_password():
 d=request.get_json(silent=True) or {};old=str(d.get('current_password',''));new=str(d.get('new_password',''))
 if len(new)<8:return jsonify(error='New password must be at least 8 characters.'),400
 a=admin_account()
 if not a or not check_password_hash(a['password_hash'],old):return jsonify(error='Current password is incorrect.'),401
 update_where('admin_accounts',{'username':f'eq.{ADMIN_USERNAME}'},{'password_hash':generate_password_hash(new),'updated_at':datetime.now(timezone.utc).isoformat()});return jsonify(ok=True,message='Password changed successfully.')
@app.post('/api/logout')
def logout():session.clear();return jsonify(ok=True)
@app.get('/api/session')
def sess():return jsonify(logged_in=session.get('role')=='admin',user=session.get('user','User'),role=session.get('role','user'))
@app.get('/api/fixtures')
def fixtures():
 s=request.args.get('sheet','').strip();f=request.args.get('fixture','').strip()
 if s not in ('AFS','TOP HAT','UNISHELL') or not f:return jsonify(items=[])
 rows=get('fixtures',{'select':'*','sheet':f'eq.{s}','fixture_no':f'eq.{f}','order':'id.asc'});b=buildable(rows)
 for x in rows:x['balance']=balance(x);x['buildable']=b
 return jsonify(items=rows)
@app.get('/api/fixture-numbers')
def fixture_numbers():
 s=request.args.get('sheet','').strip()
 if s not in ('AFS','TOP HAT','UNISHELL'):return jsonify(fixtures=[])
 rows=get('fixtures',{'select':'fixture_no','sheet':f'eq.{s}','order':'fixture_no.asc'});return jsonify(fixtures=sorted({str(x.get('fixture_no','')).strip() for x in rows if str(x.get('fixture_no','')).strip()}))
@app.post('/api/receive')
def receive():
 d=request.get_json(silent=True) or {};i=d.get('id');act=d.get('action')
 try:q=float(d.get('qty') or 0)
 except Exception:q=0
 if not i or q<=0 or act not in ('add','remove'):return jsonify(error='Enter a valid quantity'),400
 rows=get('fixtures',{'select':'*','id':f'eq.{i}'})
 if not rows:return jsonify(error='Item not found'),404
 x=rows[0];n=float(x.get('received') or 0)+q if act=='add' else max(0,float(x.get('received') or 0)-q);u=update_where('fixtures',{'id':f'eq.{i}'},{'received':n})[0]
 fr=get('fixtures',{'select':'*','sheet':f"eq.{u['sheet']}",'fixture_no':f"eq.{u['fixture_no']}"});return jsonify(ok=True,received=n,balance=balance(u),buildable=buildable(fr))
@app.post('/api/admin/update')
@require_admin
def admin_update():
 d=request.get_json(silent=True) or {};i=d.get('id');keys=['fixture_no','item_no','description','material','bom_qty','no_of_sets','total_qty','received','status'];p={k:d[k] for k in keys if k in d}
 if not i or not p:return jsonify(error='Nothing to update'),400
 for k in ('bom_qty','no_of_sets','total_qty','received'):
  if k in p:
   try:p[k]=float(p[k] or 0)
   except Exception:p[k]=0
 if 'received' in p:p['received']=max(0,p['received'])
 if 'no_of_sets' in p or 'bom_qty' in p:
  bom=float(p.get('bom_qty') or 0);sets=float(p.get('no_of_sets') or 0);p['total_qty']=max(0,bom*sets)
 if 'total_qty' in p:p['total_qty']=max(0,float(p['total_qty']))
 update_where('fixtures',{'id':f'eq.{i}'},p);return jsonify(ok=True,total_qty=p.get('total_qty'))
@app.post('/api/admin/add')
@require_admin
def admin_add():
 d=request.get_json(silent=True) or {}
 if any(not str(d.get(k,'')).strip() for k in ('sheet','fixture_no','item_no')):return jsonify(error='Sheet, Fixture No and Item No are required'),400
 bom=float(d.get('bom_qty') or 0);sets=float(d.get('no_of_sets') or 0);x={'sheet':d['sheet'],'fixture_no':str(d['fixture_no']).strip(),'item_no':str(d['item_no']).strip(),'description':d.get('description',''),'material':d.get('material',''),'bom_qty':bom,'no_of_sets':sets,'total_qty':bom*sets,'received':max(0,float(d.get('received') or 0)),'status':d.get('status','')};insert('fixtures',x);return jsonify(ok=True)
@app.delete('/api/admin/delete/<int:i>')
@require_admin
def admin_delete(i):delete_where('fixtures',{'id':f'eq.{i}'});return jsonify(ok=True)
@app.get('/api/export')
def export():
 path=os.path.join(os.path.dirname(__file__),'CORVA_TRACKER_TEMPLATE.xlsx');wb=load_workbook(path);rows=get('fixtures',{'select':'*','order':'id.asc'});by={(str(x['sheet']),str(x['fixture_no']).strip(),str(x['item_no']).strip()):x for x in rows};groups={}
 for x in rows:groups.setdefault((str(x['sheet']),str(x['fixture_no']).strip()),[]).append(x)
 fb={k:buildable(v) for k,v in groups.items()}
 for ws in wb.worksheets:
  if ws.title not in ('AFS','TOP HAT','UNISHELL'):continue
  last=None;tc=8 if ws.title=='AFS' else 7;rc=9 if ws.title=='AFS' else 8;bc=10 if ws.title=='AFS' else 9;fc=11 if ws.title=='AFS' else 10;sc=12 if ws.title=='AFS' else 11
  for r in range(4,ws.max_row+1):
   if ws.cell(r,2).value not in (None,''):last=str(ws.cell(r,2).value).strip()
   iv=ws.cell(r,3).value
   if iv in (None,'') or last is None:continue
   x=by.get((ws.title,last,str(iv).strip()))
   if not x:continue
   ws.cell(r,2).value=x['fixture_no'];ws.cell(r,3).value=x['item_no'];ws.cell(r,4).value=x.get('description');ws.cell(r,5).value=x.get('bom_qty');ws.cell(r,6).value=x.get('no_of_sets');ws.cell(r,tc).value=x.get('total_qty');ws.cell(r,rc).value=x.get('received');ws.cell(r,bc).value=balance(x);ws.cell(r,fc).value=fb.get((ws.title,last),0);ws.cell(r,sc).value=x.get('status') or ''
 out=io.BytesIO();wb.save(out);out.seek(0);return send_file(out,as_attachment=True,download_name='CORVA_FIXTURE_TRACKER_UPDATED.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
@app.get('/health')
def health():return 'OK'
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.getenv('PORT',5000)))